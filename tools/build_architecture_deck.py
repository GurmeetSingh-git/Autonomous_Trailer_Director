from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "Autonomous_Trailer_Director_Architecture_v2.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(247, 247, 242)
INK = RGBColor(18, 35, 32)
MUTED = RGBColor(86, 108, 103)
TEAL = RGBColor(26, 105, 94)
DEEP = RGBColor(11, 51, 47)
CORAL = RGBColor(241, 111, 91)
GOLD = RGBColor(222, 174, 65)
PALE = RGBColor(224, 235, 228)
WHITE = RGBColor(255, 255, 255)


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=14, color=INK, bold=False, font="Aptos", align=None, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.text = value
    p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    if align is not None: p.alignment = align
    return box


def bullet_block(slide, items, x, y, w, h, size=15, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = color
        p.level = 0; p.space_after = Pt(9)
    return box


def base_slide(kicker, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 13.333, 0.12, CORAL)
    text(slide, f"AUTONOMOUS TRAILER DIRECTOR  /  {number:02d}", 0.65, 0.42, 5.5, 0.25, 10, CORAL, True, "Aptos Mono")
    text(slide, kicker.upper(), 10.0, 0.42, 2.65, 0.25, 9, MUTED, True, "Aptos Mono", PP_ALIGN.RIGHT)
    text(slide, title, 0.65, 0.82, 11.8, 0.72, 27, INK, True, "Aptos Display")
    text(slide, str(number).zfill(2), 12.15, 6.85, 0.55, 0.25, 11, MUTED, True, "Aptos Mono", PP_ALIGN.RIGHT)
    return slide


def node(slide, label, detail, x, y, w=2.15, color=TEAL):
    rect(slide, x, y, w, 1.05, color, color, True)
    text(slide, label, x+0.14, y+0.16, w-0.28, 0.27, 14, WHITE, True)
    text(slide, detail, x+0.14, y+0.5, w-0.28, 0.38, 9, PALE)


def arrow(slide, x1, y1, x2, y2, color=CORAL):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(2.2); line.line.end_arrowhead = True


def code_tag(slide, value, x, y, w):
    rect(slide, x, y, w, 0.36, PALE, PALE, True)
    text(slide, value, x+0.12, y+0.07, w-0.24, 0.2, 10, DEEP, True, "Aptos Mono")

# 1
s = base_slide("Runtime architecture", "From uploaded video to validated trailer render", 1)
text(s, "The actual execution path", 0.72, 1.68, 4.0, 0.3, 15, CORAL, True)
text(s, "Every stage produces a concrete artifact consumed by the next stage.", 0.72, 2.02, 6.5, 0.35, 19, INK, True, "Aptos Display")
flow = [
    ("01  UPLOAD", "video.mp4", "FastAPI", DEEP),
    ("02  DEMUX", "audio.wav", "FFmpeg", TEAL),
    ("03  TRANSCRIBE", "audio_events", "Gemini / Whisper", CORAL),
    ("04  CUT DETECT", "scene cuts", "PySceneDetect", GOLD),
    ("05  BUILD", "ShotRecords", "Python", DEEP),
    ("06  VISION", "visual fields", "FFmpeg + Gemini", TEAL),
    ("07  STORY MAP", "SceneAnalysis", "structured output", CORAL),
    ("08  PLAN", "audience EDL", "planner + Validator", GOLD),
]
for i, (label, artifact, tool, color) in enumerate(flow):
    row, col = divmod(i, 4)
    x = 0.72 + col * 3.08
    y = 2.78 + row * 1.65
    node(s, label, f"{artifact}\n{tool}", x, y, 2.55, color)
    if col < 3: arrow(s, x + 2.57, y + 0.52, x + 3.0, y + 0.52)
    elif row == 0: arrow(s, x + 1.28, y + 1.08, x + 1.28, y + 1.55)
text(s, "09  RENDER  →  FFmpeg filtergraph (or OpenCV fallback)  →  trailer.mp4", 0.9, 6.25, 11.6, 0.35, 16, DEEP, True, "Aptos Mono", PP_ALIGN.CENTER)

# 2
s = base_slide("Ingestion detail", "Video is split into audio evidence and visual boundaries before analysis", 2)
rect(s, 0.8, 1.72, 2.4, 3.9, DEEP, DEEP, True)
text(s, "RAW MEDIA", 1.15, 2.1, 1.7, 0.3, 14, PALE, True, "Aptos Mono")
text(s, "episode.mp4", 1.15, 2.7, 1.7, 0.4, 22, WHITE, True, "Aptos Display")
text(s, "one source file\nmultiple streams\nunknown structure", 1.15, 3.55, 1.75, 1.0, 15, PALE)
node(s, "FFmpeg", "-vn\npcm_s16le\n16 kHz mono", 4.0, 1.95, 2.15, TEAL)
node(s, "PySceneDetect", "hard visual\ncut boundaries", 4.0, 4.05, 2.15, CORAL)
arrow(s, 3.25, 2.7, 3.92, 2.7); arrow(s, 3.25, 3.95, 3.92, 4.45)
rect(s, 7.3, 1.72, 5.0, 3.9, WHITE, PALE, True)
text(s, "TWO GROUNDED OUTPUTS", 7.7, 2.05, 4.0, 0.28, 13, CORAL, True, "Aptos Mono")
code_tag(s, "audio.wav", 7.7, 2.65, 1.65)
text(s, "transcribe_audio() → dialogue + events + timestamps", 9.55, 2.72, 2.25, 0.45, 14, INK, True)
code_tag(s, "cuts[]", 7.7, 3.55, 1.65)
text(s, "detect_scene_cuts() → start/end boundaries", 9.55, 3.62, 2.25, 0.45, 14, INK, True)
code_tag(s, "scene map", 7.7, 4.45, 1.65)
text(s, "snap_scene_to_cuts(max_drift=2.0)", 9.55, 4.52, 2.25, 0.45, 14, INK, True)
text(s, "The LLM sees structured evidence, not an unbounded video stream.", 1.0, 6.18, 11.2, 0.35, 16, MUTED, True, "Aptos Display", PP_ALIGN.CENTER)

# 3
s = base_slide("Atomic primitive", "ShotRecord: one immutable unit for audio, visual state, and coverage", 3)
rect(s, 0.8, 1.8, 5.0, 4.35, DEEP, DEEP, True)
text(s, "ShotRecord", 1.15, 2.15, 3.3, 0.35, 24, WHITE, True, "Aptos Display")
code_tag(s, "shot_id | start | end", 1.15, 2.8, 2.65)
code_tag(s, "audio: ShotAudio", 1.15, 3.35, 2.65)
code_tag(s, "visual: ShotVisual", 1.15, 3.9, 2.65)
code_tag(s, "coverage: audio_only", 1.15, 4.45, 2.65)
code_tag(s, "visual_needed: bool", 1.15, 5.0, 2.65)
text(s, "PySceneDetect owns start/end. Code attaches overlapping audio. Vision fills only confirmable visual fields.", 4.25, 3.0, 1.2, 2.1, 13, PALE)
text(s, "Architectural rule", 6.55, 2.0, 2.5, 0.3, 13, CORAL, True)
text(s, "The LLM never authors the clock.", 6.55, 2.45, 5.45, 0.65, 28, INK, True, "Aptos Display")
bullet_block(s, ["Audio: dialogue and non-verbal events", "Visual: characters, actions, location, objects", "Coverage makes uncertainty explicit instead of hiding it"], 6.6, 3.45, 5.3, 1.8, 15, MUTED)

# 4
s = base_slide("Vision optimization", "Selective frames and silent-run batching reduce multimodal cost", 4)
text(s, "Only ambiguous shots cross the vision boundary.", 0.75, 1.78, 5.3, 0.35, 20, INK, True)
node(s, "FLAG", "visual_needed", 0.85, 2.65, 1.6, CORAL)
node(s, "SAMPLE", "2–3 frames", 3.0, 2.65, 1.6, TEAL)
node(s, "GROUP", "silent runs", 5.15, 2.65, 1.6, GOLD)
node(s, "PATCH", "ShotVisual", 7.3, 2.65, 1.6, DEEP)
arrow(s, 2.5, 3.17, 2.93, 3.17); arrow(s, 4.65, 3.17, 5.08, 3.17); arrow(s, 6.8, 3.17, 7.23, 3.17)
rect(s, 9.45, 1.95, 2.9, 3.2, WHITE, PALE, True)
text(s, "Grouping rule", 9.8, 2.3, 2.2, 0.3, 16, INK, True)
text(s, "Silent + flagged shots may share one request. Dialogue-bearing shots remain singleton groups.", 9.8, 2.85, 2.1, 1.1, 15, MUTED)
code_tag(s, "group_shots_for_vision()", 9.8, 4.25, 2.1)
text(s, "Seek failures retry at timestamp − 0.5s. The pass is targeted, cacheable, and re-checkable.", 0.8, 5.45, 11.0, 0.45, 14, MUTED)

# 5
s = base_slide("Narrative aggregation", "From fine-grained evidence to macro SceneAnalysis", 5)
rect(s, 0.8, 1.85, 3.25, 3.9, DEEP, DEEP, True)
text(s, "SHOT RECORDS", 1.15, 2.25, 2.5, 0.3, 14, PALE, True, "Aptos Mono")
text(s, "50+ timestamped units", 1.15, 2.8, 2.5, 0.4, 21, WHITE, True, "Aptos Display")
text(s, "audio authority\nvisual authority\ncoverage state", 1.15, 3.65, 2.4, 1.2, 16, PALE)
rect(s, 5.0, 1.85, 3.25, 3.9, TEAL, TEAL, True)
text(s, "SCENE MAP", 5.35, 2.25, 2.5, 0.3, 14, PALE, True, "Aptos Mono")
text(s, "4–8 macro beats", 5.35, 2.8, 2.5, 0.4, 21, WHITE, True, "Aptos Display")
text(s, "descriptions\ncharacters\nemotion + spoiler risk", 5.35, 3.65, 2.4, 1.2, 16, PALE)
rect(s, 9.2, 1.85, 3.25, 3.9, CORAL, CORAL, True)
text(s, "EDIT WINDOWS", 9.55, 2.25, 2.5, 0.3, 14, WHITE, True, "Aptos Mono")
text(s, "10–15 second beats", 9.55, 2.8, 2.5, 0.4, 21, WHITE, True, "Aptos Display")
text(s, "content_start\ncontent_end\nsource_shots", 9.55, 3.65, 2.4, 1.2, 16, WHITE)
arrow(s, 4.2, 3.8, 4.9, 3.8); arrow(s, 8.4, 3.8, 9.1, 3.8)
text(s, "Conflicting modalities are preserved as evidence, not silently resolved by the model.", 1.0, 6.2, 11.0, 0.4, 15, MUTED, True)

# 6
s = base_slide("Safety gating", "Model proposals pass through deterministic policy interlocks", 6)
headers = [(0.8, "CHECK"), (4.0, "MECHANISM"), (8.0, "OUTCOME")]
for x, h in headers: text(s, h, x, 1.72, 2.3, 0.25, 11, CORAL, True, "Aptos Mono")
rows = [
    ("Existence + timing", "scene_ids; start < end", "reject / replan"),
    ("Spoiler", "protected_facts + scene evidence", "drop candidate"),
    ("Rights", "cleared_scene_ids + expiry", "warning or exclusion"),
    ("Rating + accessibility", "audience-specific checks", "warning / reject"),
    ("Grounding + injection", "evidence and content scans", "fail closed"),
]
for i, (a,b,c) in enumerate(rows):
    y = 2.15 + i*0.72
    rect(s, 0.8, y, 11.7, 0.58, WHITE if i%2==0 else PALE, WHITE if i%2==0 else PALE)
    text(s, a, 1.0, y+0.15, 2.55, 0.25, 14, INK, True)
    text(s, b, 4.0, y+0.15, 3.5, 0.25, 13, MUTED)
    text(s, c, 8.0, y+0.15, 3.6, 0.25, 13, TEAL, True)
code_tag(s, "Validator.validate() -> PASS | PASS_WITH_WARNINGS | REJECT", 3.1, 6.1, 6.9)

# 7
s = base_slide("Editorial control", "Pacing, text cards, and transitions are programmatic", 7)
text(s, "The planner produces an auditable timeline. The renderer executes it literally.", 0.75, 1.75, 7.3, 0.35, 18, INK, True)
for i, (label, detail, color) in enumerate([("hard_cut", "direct", DEEP), ("micropause", "0.3s black", TEAL), ("fade_to_black", "0.6s fade", CORAL), ("crossfade", "xfade + acrossfade", GOLD)]):
    x = 0.85 + (i%2)*3.35; y = 2.45 + (i//2)*1.25
    rect(s, x, y, 2.85, 0.9, color, color, True); text(s, label, x+0.18, y+0.17, 2.45, 0.22, 14, WHITE, True, "Aptos Mono"); text(s, detail, x+0.18, y+0.5, 2.45, 0.2, 11, PALE)
rect(s, 7.65, 2.3, 4.65, 3.35, WHITE, PALE, True)
text(s, "Text card contract", 8.0, 2.65, 3.8, 0.3, 18, INK, True)
bullet_block(s, ["3–8 grounded words", "duration clamped: 1.2–3.5s", "synthetic black clip matches source resolution/fps", "parallel transition bookkeeping prevents card index drift"], 8.0, 3.25, 3.8, 1.8, 14, MUTED)

# 8
s = base_slide("Human-in-the-loop", "Checkpointed runs make failure recoverable and feedback actionable", 8)
node(s, "ANALYZE", "expensive pass", 0.9, 2.15, 2.0, DEEP)
node(s, "CHECKPOINT", "runtime/pending", 3.45, 2.15, 2.0, GOLD)
node(s, "DIRECTOR", "promise + feedback", 6.0, 2.15, 2.0, CORAL)
node(s, "REPLAN", "review graph", 8.55, 2.15, 2.0, TEAL)
arrow(s, 2.95, 2.68, 3.4, 2.68); arrow(s, 5.5, 2.68, 5.95, 2.68); arrow(s, 8.05, 2.68, 8.5, 2.68)
rect(s, 1.0, 4.25, 10.9, 1.35, WHITE, PALE, True)
text(s, "Persisted evidence", 1.35, 4.55, 2.0, 0.25, 15, CORAL, True)
text(s, "cuts  •  transcripts  •  ShotRecords  •  source media  •  selected model", 3.4, 4.55, 7.8, 0.25, 16, INK, True)
text(s, "A story-map failure retries the model call without repeating extraction, transcription, or vision analysis.", 1.35, 5.0, 9.8, 0.3, 14, MUTED)

# 9
s = base_slide("Rendering", "Deterministic NLE assembly with an OpenCV fallback", 9)
rect(s, 0.8, 1.85, 5.45, 3.85, DEEP, DEEP, True)
text(s, "FFmpeg primary", 1.2, 2.25, 3.2, 0.3, 22, WHITE, True, "Aptos Display")
bullet_block(s, ["AAC stereo at 48 kHz", "YUV420p output", "avoid_negative_ts make_zero", "concat list + xfade/acrossfade", "text cards and black holds"], 1.2, 2.95, 4.3, 2.0, 15, PALE)
rect(s, 7.0, 1.85, 5.45, 3.85, TEAL, TEAL, True)
text(s, "OpenCV fallback", 7.4, 2.25, 3.4, 0.3, 22, WHITE, True, "Aptos Display")
bullet_block(s, ["activates when FFmpeg is absent", "frame-level seek via CAP_PROP_POS_FRAMES", "VideoWriter mp4v output", "keeps local replay usable"], 7.4, 2.95, 4.3, 1.8, 15, PALE)
text(s, "The output is assembled from verified source ranges, never generated as a new video by the model.", 1.0, 6.2, 11.0, 0.4, 15, MUTED, True)

# 10
s = base_slide("Technical highlights", "A replayable foundation for production-grade editorial automation", 10)
text(s, "What this architecture buys us", 0.8, 1.8, 4.0, 0.35, 20, INK, True)
items = [("Determinism", "same evidence -> same replay plan", CORAL), ("Cost control", "vision only where ambiguity earns it", TEAL), ("Fault tolerance", "checkpoint + fallback renderer", GOLD), ("Auditability", "decisions, evidence, validation artifacts", DEEP)]
for i, (label, detail, color) in enumerate(items):
    y = 2.45 + i*0.78
    rect(s, 0.85, y, 0.18, 0.5, color, color)
    text(s, label, 1.25, y+0.04, 2.0, 0.25, 16, INK, True)
    text(s, detail, 3.4, y+0.04, 4.4, 0.25, 15, MUTED)
rect(s, 8.35, 1.85, 3.95, 3.95, DEEP, DEEP, True)
text(s, "Next extensions", 8.75, 2.25, 3.0, 0.3, 20, WHITE, True, "Aptos Display")
bullet_block(s, ["beat-level rhythm detection", "rights and contract adapters", "renderer observability", "human approval analytics", "multi-model semantic routing"], 8.75, 3.0, 2.9, 1.9, 15, PALE)
text(s, "The boundary stays stable: extract deterministically, reason sparsely, validate independently, render reproducibly.", 0.8, 6.25, 11.4, 0.45, 17, CORAL, True)

prs.save(OUT)
print(OUT)
